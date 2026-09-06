"""
NCAA RPI for Division I women's soccer.

history(registry)  End-of-season RPI 2007-2024 from Chris Henderson's "RPI for Division I Women's
                   Soccer" Google Sheets (Teams tab, CSV export). -> public/data/rpi/<year>.json
current(registry)  Latest weekly RPI table from ncaa.com (server-rendered). -> public/data/rpi/current.json
                   and an immutable snapshot public/data/rpi/weekly/<season>/<through-date>.json so the
                   repo accumulates the in-season history the NCAA does not publish.

Team names differ between the two sources (Henderson: 'NorthCarolinaU', NCAA: 'North Carolina');
build.py joins them to programs via registry ids.rpiHistoryName / ids.ncaaName.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import os
import re

from bs4 import BeautifulSoup

from . import common

HIST_COLS = {
    "team": "Team", "region": "Region", "conference": "Conference",
    "rpiRank": "NCAA RPI Rank", "sosRank": "NCAA Strength of Schedule Contributor Rank",
    "oppAvgRpiRank": "Opponents Average NCAA RPI Rank",
    "nonConfOppAvgRpiRank": "Non Conference Opponents Average NCAA RPI Rank",
    "top50ResultsRank": "NCAA RPI Top 50 Results Rank",
    "balancedRpiRank": "Balanced RPI Rank", "kpiRank": "KPI Rank", "masseyRank": "Massey Rank",
    "ncaaSeed": "NCAA Tournament Actual Seed or Selection",
    "autoQualifier": "NCAA Tournament Automatic Qualifier",
}
MONTHS = {m: i + 1 for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}


def _int(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def parse_history_csv(text: str) -> list[dict]:
    rows = list(csv.reader(io.StringIO(text)))
    return parse_history_rows(rows)


def parse_history_xlsx(blob: bytes) -> list[dict]:
    """The 2021-2024 sheets compute the Team column with formulas that the CSV export renders as
    #ERROR!, but the XLSX export carries cached values. Needs openpyxl."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
    if "Teams" not in wb.sheetnames:
        return []
    rows = [["" if v is None else str(v) for v in r] for r in wb["Teams"].iter_rows(values_only=True)]
    return parse_history_rows(rows)


def parse_history_rows(rows: list[list[str]]) -> list[dict]:
    if not rows:
        return []
    hdr = [h.strip() for h in rows[0]]
    # Duplicate header names exist; take the LAST occurrence for seed/AQ (the 'actual' block),
    # first occurrence for everything else.
    idx = {}
    for key, col in HIST_COLS.items():
        hits = [i for i, h in enumerate(hdr) if h.strip() == col]
        if not hits:
            continue
        idx[key] = hits[-1] if key in ("ncaaSeed", "autoQualifier") else hits[0]
    if "team" not in idx or "rpiRank" not in idx:
        return []
    out = []
    for r in rows[1:]:
        if len(r) <= idx["team"]:
            continue
        team = r[idx["team"]].strip()
        if not team or team.startswith("#"):
            continue
        rec = {"team": team}
        for key, i in idx.items():
            if key == "team":
                continue
            v = r[i].strip() if i < len(r) else ""
            rec[key] = v if key in ("region", "conference") else _int(v)
        if rec.get("rpiRank") is None:
            continue
        out.append(rec)
    out.sort(key=lambda x: x["rpiRank"])
    return out


def history(registry: dict, years: list[int] | None = None, *, force: bool = False) -> dict:
    src = registry["sources"]["rpiHistory"]
    sheets = src["sheets"]
    years = years or sorted(int(y) for y in sheets)
    summary = {}
    for y in years:
        sheet_id = sheets.get(str(y))
        if not sheet_id:
            continue
        out_path = os.path.join(common.RPI_OUT_DIR, f"{y}.json")
        if os.path.exists(out_path) and not force:
            existing = common.read_json(out_path, {})
            summary[y] = f"kept ({len(existing.get('teams', []))} teams)"
            continue
        teams, meta = [], {}
        # XLSX first (cached formula values), CSV as fallback.
        for kind, url in (("xlsx", src["xlsxTemplate"].format(sheetId=sheet_id)),
                          ("csv", src["csvTemplate"].format(sheetId=sheet_id))):
            try:
                blob, meta = common.fetch(url, max_age_hours=24 * 365, timeout=180)
                teams = parse_history_xlsx(blob) if kind == "xlsx" else parse_history_csv(blob.decode("utf-8", "replace"))
            except ImportError:
                common.log("rpi history: openpyxl not installed (pip install openpyxl); trying CSV")
                continue
            except common.FetchError as e:
                common.log(f"rpi history {y}: {kind} fetch failed: {e}")
                continue
            if teams:
                break
        if not teams:
            summary[y] = "no usable rows"
            common.log(f"rpi history {y}: {summary[y]}")
            continue
        common.write_json(out_path, {
            "year": y, "kind": "end-of-season",
            "provider": src.get("provider", "RPI for Division I Women's Soccer (Chris Henderson)"),
            "sourceUrl": f"https://docs.google.com/spreadsheets/d/{sheet_id}/",
            "note": "Ranks recomputed by the provider under the 2024 NCAA RPI formula; may differ slightly from the NCAA's published ranks of that season.",
            "fetchedAt": meta.get("fetchedAt"), "teams": teams,
        })
        summary[y] = f"{len(teams)} teams"
        common.log(f"rpi history {y}: {len(teams)} teams (#1 {teams[0]['team']})")
    common.update_refresh_state("rpi.history", {"years": {str(k): v for k, v in summary.items()}})
    return summary


def _parse_through(text: str) -> str | None:
    m = re.search(r"Through Games\s+([A-Za-z]{3})\.?\s+(\d{1,2}),?\s+(\d{4})", text)
    if not m:
        return None
    mon = MONTHS.get(m.group(1).lower()[:3])
    return f"{m.group(3)}-{mon:02d}-{int(m.group(2)):02d}" if mon else None


def parse_current_html(html: str) -> dict:
    soup = BeautifulSoup(html, "html.parser")
    fig = soup.select_one(".rankings-last-updated")
    through = _parse_through(fig.get_text(" ")) if fig else _parse_through(soup.get_text(" "))
    table = soup.select_one("table.sticky") or soup.find("table")
    teams = []
    if table:
        heads = [common.clean(th.get_text()).lower() for th in table.select("thead th")]
        for tr in table.select("tbody tr"):
            cells = [common.clean(td.get_text()) for td in tr.find_all("td")]
            if len(cells) < 3:
                continue
            row = dict(zip(heads, cells))
            rank = _int(row.get("rank"))
            if rank is None:
                continue
            teams.append({
                "rank": rank, "school": row.get("school"), "record": row.get("record"),
                "conference": row.get("conf") or row.get("conference"),
                "road": row.get("road"), "neutral": row.get("neutral"), "home": row.get("home"),
                "nonDiv1": row.get("non-div i"), "prevRank": _int(row.get("prev")),
            })
    return {"throughGames": through, "teams": teams}


def current(registry: dict) -> dict:
    url = registry["sources"]["ncaaRpi"]["current"]
    html, meta = common.fetch_text(url, max_age_hours=6)
    parsed = parse_current_html(html)
    if len(parsed["teams"]) < 100:
        raise common.FetchError(f"rpi current: only {len(parsed['teams'])} rows parsed from {url}")
    through = parsed["throughGames"]
    season = int(through[:4]) if through else dt.date.today().year
    data = {"kind": "weekly", "season": season, "throughGames": through, "sourceUrl": url,
            "fetchedAt": meta.get("fetchedAt") or common.now_iso(), "teams": parsed["teams"]}
    common.write_json(os.path.join(common.RPI_OUT_DIR, "current.json"), data)
    if through:
        snap = os.path.join(common.RPI_OUT_DIR, "weekly", str(season), f"{through}.json")
        if not os.path.exists(snap):
            common.write_json(snap, data)
            common.log(f"rpi current: new snapshot {season}/{through} ({len(parsed['teams'])} teams)")
        else:
            common.log(f"rpi current: unchanged (through {through})")
    common.update_refresh_state("rpi.current", {"throughGames": through, "teams": len(parsed["teams"])})
    return data


def weekly_index() -> dict:
    """List archived weekly snapshots: {season: [dates]}"""
    base = os.path.join(common.RPI_OUT_DIR, "weekly")
    out = {}
    if not os.path.isdir(base):
        return out
    for season in sorted(os.listdir(base)):
        files = sorted(f[:-5] for f in os.listdir(os.path.join(base, season)) if f.endswith(".json"))
        out[season] = files
    return out
